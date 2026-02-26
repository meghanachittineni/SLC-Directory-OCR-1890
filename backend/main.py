import os
import io
import base64
import json
import re
import tempfile
import traceback
from pathlib import Path
from typing import Optional

import google.generativeai as genai
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from PIL import Image
import pdf2image

load_dotenv()

app = FastAPI(title="Salt Lake City Directory OCR API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
if not GEMINI_API_KEY:
    raise RuntimeError("GEMINI_API_KEY not set in .env")

genai.configure(api_key=GEMINI_API_KEY)

SYSTEM_PROMPT = """You are an expert OCR system specializing in 19th century city directories.
You will receive an image of a page from the 1890 Salt Lake City Directory.

Extract EVERY person listed on the page and return them as a JSON array.
Each person should have these fields:
- last_name: string
- first_name: string (include titles like Miss, Mrs, etc. in parentheses if present)
- occupation: string (expand abbreviations: clk=clerk (clk), carp=carpenter (carp), res=residence, bds=boards, rms=rooms, wid=widow, lab=laborer, mach=machine/machinery or machinist, mkr=maker, mnfr=manufacturer, agt=agent, asst=assistant, opr=operator, propr=proprietor, supt=superintendent, treas=treasurer, sec=secretary, cor=corner, e=east, w=west, n=north, s=south) - use null if no occupation listed
- company: string or null (company or employer if listed)
- business_address: string or null
- status: string (use "residence (res)", "boards (bds)", "rooms (rms)", "room (rm)", or null)
- address: string or null (the home/boarding address)
- flag: string or null (any special notes)

Common patterns in the directory:
- "Abbott Charles (C Abbott & Son), res 349 W Sixth South" means last=Abbott, first=Charles, company=C Abbott & Son, status=residence (res), address=349 W Sixth South
- "Aaron Cyrus, clk Dunkley & Co, bds 80 W First South" means last=Aaron, first=Cyrus, occupation=clerk (clk), company=Dunkley & Co, status=boards (bds), address=80 W First South
- "Abbott Emily, machine hand Solomon Bros, bds 349 W Fifth South" means occupation=machine hand, company=Solomon Bros
- If a person is listed with "(C Abbott & Son)" after their name, they are owner/partner: occupation=[Owner/Partner]
- Owners/Partners of businesses: use occupation="[Owner/Partner]"
- Widow entries: include "wid" in name context

Return ONLY a valid JSON array, no markdown, no explanation.
Example:
[
  {
    "last_name": "Aaron",
    "first_name": "Cyrus",
    "occupation": "clerk (clk)",
    "company": "Dunkley & Co",
    "business_address": null,
    "status": "boards (bds)",
    "address": "80 W First South",
    "flag": null
  }
]
"""


def image_to_base64(img: Image.Image) -> str:
    """Convert PIL image to base64 string."""
    buffered = io.BytesIO()
    img.save(buffered, format="JPEG", quality=95)
    return base64.b64encode(buffered.getvalue()).decode("utf-8")


def extract_images_from_pdf(pdf_bytes: bytes) -> list[Image.Image]:
    """Convert PDF pages to images."""
    images = pdf2image.convert_from_bytes(pdf_bytes, dpi=300)
    return images


def ocr_image_with_gemini(img: Image.Image) -> list[dict]:
    """Send image to Gemini and get structured data back."""
    model = genai.GenerativeModel("gemini-2.5-flash")

    img_data = {
        "mime_type": "image/jpeg",
        "data": image_to_base64(img),
    }

    response = model.generate_content(
        [SYSTEM_PROMPT, img_data],
        generation_config=genai.types.GenerationConfig(
            temperature=0.1,
            max_output_tokens=8192,
        ),
    )

    text = response.text.strip()

    # Strip markdown code fences if present
    text = re.sub(r"^```json\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    text = text.strip()

    data = json.loads(text)
    return data


def build_excel(records: list[dict]) -> bytes:
    """Build a styled Excel file from extracted records."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Directory"

    headers = [
        "Last Name", "First Name", "Occupation", "Company / Affiliation",
        "Business Address", "Status", "Address", "Flag"
    ]

    # Header style
    header_fill = PatternFill(start_color="1B3A5C", end_color="1B3A5C", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border

    ws.row_dimensions[1].height = 25

    # Alternate row colors
    fill_light = PatternFill(start_color="EEF3F8", end_color="EEF3F8", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="Calibri", size=10)
    data_align = Alignment(vertical="center", wrap_text=True)

    field_keys = [
        "last_name", "first_name", "occupation", "company",
        "business_address", "status", "address", "flag"
    ]

    for row_idx, record in enumerate(records, 2):
        fill = fill_light if row_idx % 2 == 0 else fill_white
        for col_idx, key in enumerate(field_keys, 1):
            val = record.get(key)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = fill
            cell.font = data_font
            cell.alignment = data_align
            cell.border = border

    # Column widths
    col_widths = [15, 20, 25, 30, 25, 20, 35, 15]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:H{len(records) + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/ocr")
async def ocr_directory(file: UploadFile = File(...)):
    """Process uploaded PDF or image and return structured directory data."""
    filename = file.filename or ""
    content = await file.read()
    ext = Path(filename).suffix.lower()

    try:
        if ext == ".pdf":
            images = extract_images_from_pdf(content)
        elif ext in [".jpg", ".jpeg", ".png", ".tiff", ".tif", ".webp"]:
            img = Image.open(io.BytesIO(content))
            images = [img]
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}")

        all_records = []
        errors = []

        for page_num, img in enumerate(images, 1):
            try:
                records = ocr_image_with_gemini(img)
                all_records.extend(records)
            except Exception as e:
                errors.append(f"Page {page_num}: {str(e)}")

        if not all_records and errors:
            raise HTTPException(status_code=500, detail=f"OCR failed: {'; '.join(errors)}")

        return {
            "records": all_records,
            "total": len(all_records),
            "pages_processed": len(images),
            "errors": errors,
        }

    except HTTPException:
        raise
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/export")
async def export_excel(file: UploadFile = File(...)):
    """Process file and return Excel download directly."""
    filename = file.filename or ""
    content = await file.read()
    ext = Path(filename).suffix.lower()

    try:
        if ext == ".pdf":
            images = extract_images_from_pdf(content)
        elif ext in [".jpg", ".jpeg", ".png", ".tiff", ".tif", ".webp"]:
            img = Image.open(io.BytesIO(content))
            images = [img]
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}")

        all_records = []
        for page_num, img in enumerate(images, 1):
            try:
                records = ocr_image_with_gemini(img)
                all_records.extend(records)
            except Exception as e:
                print(f"Page {page_num} error: {e}")

        excel_bytes = build_excel(all_records)

        return StreamingResponse(
            io.BytesIO(excel_bytes),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=directory_extracted.xlsx"},
        )

    except HTTPException:
        raise
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/export-from-records")
async def export_from_records(records: list[dict]):
    """Build Excel from already-extracted records."""
    excel_bytes = build_excel(records)
    return StreamingResponse(
        io.BytesIO(excel_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=directory_extracted.xlsx"},
    )
